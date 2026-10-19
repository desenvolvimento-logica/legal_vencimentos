"""Planilha de pendências da importação do sistema anterior.

Lê o JSON "pendentes_por_cliente" devolvido por legal_legacy_import.sql e as
planilhas originais (docs/*.xlsx, fora do git) e gera
docs/Pendencias_importacao.xlsx (fora do git) com as linhas completas que
não foram importadas, para cadastro manual pelas operadoras.

Uso: python -I scripts/legacy-import/pending.py <pendentes.json>   (requer openpyxl)
"""

import collections
import datetime
import json
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
from generate import DOCUMENT_TYPES, ROOT, read_rows  # noqa: E402

OUTPUT = ROOT / "docs" / "Pendencias_importacao.xlsx"

TYPE_LABELS = {"CERTIFICADO": "Certificado", "LICENCA": "Licença", "OUTRO": "Outros"}

WHAT_TO_DO = {
    "filial sem CNPJ na planilha": "Localizar a filial no Acessórias (pelo CNPJ) e cadastrar o vencimento nela.",
    "mais de um cliente ativo com esse nome": "Há matriz/filiais com o mesmo nome: conferir o CNPJ e cadastrar no cliente certo.",
    "cliente não encontrado no Acessórias": "Nome diferente no Acessórias (ou cliente fora da Legalização): localizar o cliente e cadastrar.",
    "cliente inativo no Acessórias": "Cliente inativo: confirmar se ainda é atendido; se sim, reativar no Acessórias.",
}


def main():
    if len(sys.argv) != 2:
        sys.exit("Uso: python -I scripts/legacy-import/pending.py <pendentes.json>")
    pending = json.loads(Path(sys.argv[1]).read_text(encoding="utf-8"))
    reasons = {item["cliente"]: item["motivo"] for item in pending}
    expected = {item["cliente"]: item["linhas"] for item in pending}

    rows = [r for r in read_rows() if r[2] in reasons]
    found = collections.Counter(r[2] for r in rows)
    mismatched = {n: (expected[n], found.get(n, 0)) for n in expected if expected[n] != found.get(n, 0)}
    if mismatched:
        sys.exit(f"Contagem não confere com o JSON (esperado, encontrado): {mismatched}")

    order = list(WHAT_TO_DO)
    rows.sort(key=lambda r: (order.index(reasons[r[2]]), r[2], r[5]))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Pendências"
    headers = ["Cliente (sistema anterior)", "Motivo", "O que fazer", "Documento",
               "Tipo", "Vencimento", "Observação", "Planilha", "Linha"]
    ws.append(headers)
    for sheet, line, client, document, kind, due, notes in rows:
        reason = reasons[client]
        ws.append([client, reason, WHAT_TO_DO[reason], document, TYPE_LABELS[kind],
                   datetime.date.fromisoformat(due), notes, sheet, line])

    header_fill = PatternFill("solid", fgColor="101827")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")
    widths = [48, 34, 60, 44, 12, 12, 60, 30, 8]
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    for row in ws.iter_rows(min_row=2):
        row[5].number_format = "DD/MM/YYYY"
        for cell in (row[2], row[6]):
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for cell in row:
            if cell.alignment.wrap_text is not True:
                cell.alignment = Alignment(vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    summary = wb.create_sheet("Resumo")
    summary.append(["Motivo", "Clientes", "Linhas", "O que fazer"])
    for reason in order:
        clients = sum(1 for item in pending if item["motivo"] == reason)
        lines = sum(item["linhas"] for item in pending if item["motivo"] == reason)
        summary.append([reason, clients, lines, WHAT_TO_DO[reason]])
    summary.append(["Total", len(pending), len(rows), ""])
    for cell in summary[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
    for cell in summary[summary.max_row]:
        cell.font = Font(bold=True)
    for index, width in enumerate([40, 10, 10, 90], start=1):
        summary.column_dimensions[get_column_letter(index)].width = width

    wb.save(OUTPUT)
    print(f"{len(rows)} linhas pendentes de {len(pending)} clientes -> {OUTPUT}")


if __name__ == "__main__":
    assert DOCUMENT_TYPES  # mesma tabela de tipos da importação
    main()
